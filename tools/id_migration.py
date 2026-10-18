#!/usr/bin/env python3
"""Export, apply, and audit the approved THEMELIOS ID migration map.

Only registry-keyed blockstate/item-model filenames, data-pack registry
references, and existing language keys move. Internal block-model, texture,
and published Continuity paths stay stable for resource-pack compatibility.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

EXPECTED_TOTAL = 861
EXPECTED_ALIASES = 861
EXPECTED_QUATREFOIL_ALIASES = 48
EXPECTED_CYLINDER_ALIASES = 54
EXPECTED_REASON_COUNTS = {
    "Aged order": 243,
    "Aged order; Slice order": 54,
    "Inlay grammar": 144,
    "Inlay grammar; Rose → Rosette": 48,
    "Slice order": 372,
}
MANIFEST_COLUMNS = (
    "source_row",
    "mode",
    "old_path",
    "canonical_path",
    "old_display_name",
    "canonical_display_name",
    "design_style",
    "publication_status",
    "reason",
    "review_status",
)
JSON_LINE = re.compile(
    r'^(?P<indent>\s*)"(?P<key>(?:\\.|[^"\\])+)"\s*:\s*'
    r'"(?P<value>(?:\\.|[^"\\])*)"(?P<comma>,?)\s*$'
)


@dataclass(frozen=True)
class Entry:
    source_row: int
    mode: str
    old_path: str
    canonical_path: str
    old_display_name: str
    canonical_display_name: str
    design_style: str
    publication_status: str
    reason: str
    review_status: str


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    subcommands = result.add_subparsers(dest="command", required=True)

    export = subcommands.add_parser("export", help="export approved THEMELIOS rows from a workbook copy")
    export.add_argument("--workbook", type=Path, required=True)
    export.add_argument("--output", type=Path, required=True)
    export.add_argument("--sheet", default="05 ID Migration Map")

    apply = subcommands.add_parser("apply", help="apply registry-keyed source moves")
    apply.add_argument("--manifest", type=Path, required=True)
    apply.add_argument("--repo", type=Path, default=Path.cwd())

    audit = subcommands.add_parser("audit", help="audit the committed source migration")
    audit.add_argument("--manifest", type=Path, required=True)
    audit.add_argument("--repo", type=Path, default=Path.cwd())
    audit.add_argument("--report", type=Path)
    return result


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def validate_entries(entries: list[Entry]) -> None:
    require(len(entries) == EXPECTED_TOTAL,
            f"expected {EXPECTED_TOTAL} THEMELIOS migrations, found {len(entries)}")
    require(all(entry.mode == "PERMANENT_ALIAS" for entry in entries),
            "every THEMELIOS migration must be a permanent alias")
    require(all(entry.review_status == "Approved" for entry in entries),
            "every migration row must be Approved")
    require(all(entry.publication_status == "Published — v1.5.12 source tag" for entry in entries),
            "every migration row must retain the v1.5.12 publication evidence")

    old_paths = [entry.old_path for entry in entries]
    canonical_paths = [entry.canonical_path for entry in entries]
    require(len(set(old_paths)) == len(old_paths), "duplicate old paths in migration map")
    require(len(set(canonical_paths)) == len(canonical_paths),
            "duplicate canonical paths in migration map")
    require(not set(old_paths).intersection(canonical_paths), "alias chains are forbidden")
    require(all(entry.old_path != entry.canonical_path for entry in entries),
            "migration rows must change the ID")
    require(Counter(entry.reason for entry in entries) == Counter(EXPECTED_REASON_COUNTS),
            "migration reason counts drifted from the approved workbook")
    require(sum("quatrefoil" in entry.old_path for entry in entries)
            == EXPECTED_QUATREFOIL_ALIASES,
            f"expected {EXPECTED_QUATREFOIL_ALIASES} Quatrefoil aliases")
    require(sum("cylinder_" in entry.old_path for entry in entries)
            == EXPECTED_CYLINDER_ALIASES,
            f"expected {EXPECTED_CYLINDER_ALIASES} aged cylinder aliases")
    require(not any("column_round" in entry.canonical_path for entry in entries),
            "held cylinder-to-column proposals leaked into the approved migration")


def workbook_entries(workbook: Path, sheet_name: str) -> list[Entry]:
    try:
        from openpyxl import load_workbook
    except ImportError as exception:
        raise RuntimeError("openpyxl is required only for workbook export") from exception

    worksheet = load_workbook(workbook, read_only=True, data_only=True)[sheet_name]
    headers = list(next(worksheet.iter_rows(min_row=4, max_row=4, values_only=True)))
    index = {str(value): position for position, value in enumerate(headers)}
    required = {
        "Mod", "Design style", "Current block ID", "Proposed block ID",
        "Current item ID", "Proposed item ID", "Current display name",
        "Proposed display name", "Publication status", "Migration action",
        "Block alias", "Item alias", "Reason", "Collision check", "Review status",
    }
    missing = required.difference(index)
    require(not missing, f"missing workbook columns: {sorted(missing)}")

    entries: list[Entry] = []
    for source_row, values in enumerate(
            worksheet.iter_rows(min_row=5, values_only=True), start=5):
        if str(values[index["Mod"]] or "").strip().upper() != "THEMELIOS":
            continue

        action = str(values[index["Migration action"]] or "").strip().upper()
        require(action == "PERMANENT ALIAS",
                f"row {source_row}: unsupported migration action {action!r}")
        old_block = str(values[index["Current block ID"]] or "").strip()
        new_block = str(values[index["Proposed block ID"]] or "").strip()
        old_item = str(values[index["Current item ID"]] or "").strip()
        new_item = str(values[index["Proposed item ID"]] or "").strip()
        require(old_block == old_item, f"row {source_row}: block/item old IDs differ")
        require(new_block == new_item, f"row {source_row}: block/item canonical IDs differ")
        require(old_block.startswith("themelios:") and new_block.startswith("themelios:"),
                f"row {source_row}: IDs must stay in the themelios namespace")
        require(str(values[index["Collision check"]] or "").strip() == "Clear",
                f"row {source_row}: collision check is not Clear")
        require(str(values[index["Block alias"]] or "").strip() == "Yes"
                and str(values[index["Item alias"]] or "").strip() == "Yes",
                f"row {source_row}: aliases must cover both block and item")

        entries.append(Entry(
            source_row=source_row,
            mode="PERMANENT_ALIAS",
            old_path=old_block.split(":", 1)[1],
            canonical_path=new_block.split(":", 1)[1],
            old_display_name=str(values[index["Current display name"]] or "").strip(),
            canonical_display_name=str(values[index["Proposed display name"]] or "").strip(),
            design_style=str(values[index["Design style"]] or "").strip(),
            publication_status=str(values[index["Publication status"]] or "").strip(),
            reason=str(values[index["Reason"]] or "").strip(),
            review_status=str(values[index["Review status"]] or "").strip(),
        ))

    validate_entries(entries)
    return entries


def write_manifest(entries: Iterable[Entry], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=MANIFEST_COLUMNS, delimiter="\t",
                                lineterminator="\n")
        writer.writeheader()
        for entry in entries:
            writer.writerow(entry.__dict__)


def read_manifest(path: Path) -> list[Entry]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        require(tuple(reader.fieldnames or ()) == MANIFEST_COLUMNS,
                f"unexpected manifest columns in {path}")
        entries = [Entry(source_row=int(row["source_row"]),
                         **{key: row[key] for key in MANIFEST_COLUMNS if key != "source_row"})
                   for row in reader]
    validate_entries(entries)
    return entries


def move_registry_resources(repo: Path, entries: list[Entry]) -> int:
    roots = (
        repo / "src/main/resources/assets/themelios/blockstates",
        repo / "src/main/resources/assets/themelios/models/item",
    )
    moved = 0
    for entry in entries:
        for root in roots:
            old = root / f"{entry.old_path}.json"
            canonical = root / f"{entry.canonical_path}.json"
            require(not (old.exists() and canonical.exists()),
                    f"both old and canonical resources exist: {old} / {canonical}")
            if old.exists():
                old.rename(canonical)
                moved += 1
            else:
                require(canonical.is_file(),
                        f"missing old and canonical registry resource for row {entry.source_row}: {root}")
    return moved


def update_registry_references(repo: Path, entries: list[Entry]) -> int:
    replacements = {
        f"themelios:{entry.old_path}": f"themelios:{entry.canonical_path}"
        for entry in entries
    }
    pattern = re.compile(
        r"(?:"
        + "|".join(re.escape(value) for value in sorted(replacements, key=len, reverse=True))
        + r")(?![a-z0-9_])"
    )
    changed = 0
    for path in sorted((repo / "src/main/resources/data").rglob("*.json")):
        raw = path.read_text(encoding="utf-8-sig")
        updated, count = pattern.subn(lambda match: replacements[match.group(0)], raw)
        if count:
            path.write_text(updated, encoding="utf-8", newline="")
            changed += count
        json.loads(updated)
    return changed


def update_language(path: Path, language: str, entries: list[Entry],
                    require_complete: bool) -> tuple[int, int, int]:
    raw = path.read_text(encoding="utf-8-sig")
    newline = "\r\n" if "\r\n" in raw else "\n"
    lines = raw.splitlines()
    key_to_entry = {f"block.themelios.{entry.old_path}": entry for entry in entries}
    key_to_entry.update({f"block.themelios.{entry.canonical_path}": entry for entry in entries})
    canonical_seen: set[str] = set()
    changed_keys = 0
    changed_values = 0
    output: list[str] = []

    for line in lines:
        match = JSON_LINE.match(line)
        if match is None or match.group("key") not in key_to_entry:
            output.append(line)
            continue

        entry = key_to_entry[match.group("key")]
        canonical_key = f"block.themelios.{entry.canonical_path}"
        if canonical_key in canonical_seen:
            changed_keys += 1
            continue
        canonical_seen.add(canonical_key)

        current_value = json.loads(f'"{match.group("value")}"')
        next_value = entry.canonical_display_name if language == "en_us" else current_value
        if match.group("key") != canonical_key:
            changed_keys += 1
        if current_value != next_value:
            changed_values += 1
        output.append(
            f'{match.group("indent")}{json.dumps(canonical_key, ensure_ascii=False)}: '
            f'{json.dumps(next_value, ensure_ascii=False)}{match.group("comma")}'
        )

    if require_complete:
        missing = {f"block.themelios.{entry.canonical_path}" for entry in entries}.difference(canonical_seen)
        require(not missing, f"missing canonical language keys in {path}: {sorted(missing)[:5]}")
    path.write_text(newline.join(output) + newline, encoding="utf-8", newline="")
    json.loads(path.read_text(encoding="utf-8-sig"))
    return changed_keys, changed_values, len(canonical_seen)


def apply_source(repo: Path, entries: list[Entry]) -> dict[str, int]:
    repo = repo.resolve()
    moved = move_registry_resources(repo, entries)
    changed_registry_references = update_registry_references(repo, entries)
    changed_keys = 0
    changed_values = 0
    localized_keys = 0
    for language in ("en_us", "de_de", "es_es"):
        path = repo / f"src/main/resources/assets/themelios/lang/{language}.json"
        keys, values, seen = update_language(
            path, language, entries, require_complete=True)
        changed_keys += keys
        changed_values += values
        localized_keys += seen
    return {
        "moved_registry_resources": moved,
        "changed_registry_references": changed_registry_references,
        "changed_language_keys": changed_keys,
        "changed_language_values": changed_values,
        "localized_migration_keys": localized_keys,
    }


def audit_source(repo: Path, entries: list[Entry]) -> dict[str, object]:
    repo = repo.resolve()
    errors: list[str] = []
    canonical_resources = 0
    stale_resources = 0
    for entry in entries:
        for relative_root in (
                "src/main/resources/assets/themelios/blockstates",
                "src/main/resources/assets/themelios/models/item"):
            root = repo / relative_root
            old = root / f"{entry.old_path}.json"
            canonical = root / f"{entry.canonical_path}.json"
            if canonical.is_file():
                canonical_resources += 1
            else:
                errors.append(f"missing canonical resource: {canonical}")
            if old.exists():
                stale_resources += 1
                errors.append(f"stale old registry resource: {old}")

    language_keys = 0
    stale_language_keys = 0
    language_coverage: dict[str, int] = {}
    for language in ("en_us", "de_de", "es_es"):
        path = repo / f"src/main/resources/assets/themelios/lang/{language}.json"
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        coverage = 0
        for entry in entries:
            canonical_key = f"block.themelios.{entry.canonical_path}"
            old_key = f"block.themelios.{entry.old_path}"
            if canonical_key in data:
                language_keys += 1
                coverage += 1
            else:
                errors.append(f"missing language key in {language}: {canonical_key}")
            if old_key in data:
                stale_language_keys += 1
                errors.append(f"stale language key in {language}: {old_key}")
        language_coverage[language] = coverage

    stale_registry_references = 0
    old_registry_ids = {f"themelios:{entry.old_path}" for entry in entries}
    id_pattern = re.compile(r"themelios:[a-z0-9_]+")
    for path in sorted((repo / "src/main/resources/data").rglob("*.json")):
        raw = path.read_text(encoding="utf-8-sig")
        try:
            json.loads(raw)
        except json.JSONDecodeError as exception:
            errors.append(f"invalid JSON in {path}: {exception}")
            continue
        stale = sorted(set(id_pattern.findall(raw)).intersection(old_registry_ids))
        if stale:
            stale_registry_references += sum(raw.count(value) for value in stale)
            errors.append(f"stale registry ID reference in {path}: {stale[:5]}")

    ctm_legacy_references = 0
    ctm_property_boms = 0
    ctm_roots = sorted((repo / "src/main/resources").rglob("*.properties"))
    legacy_tokens = {
        f"themelios:{entry.old_path}" for entry in entries
    } | {
        f"erydon:{entry.old_path}" for entry in entries
    }
    for path in ctm_roots:
        raw_bytes = path.read_bytes()
        if raw_bytes.startswith(b"\xef\xbb\xbf"):
            ctm_property_boms += 1
            errors.append(f"CTM property has UTF-8 BOM: {path}")
        raw = raw_bytes.decode("utf-8-sig")
        ctm_legacy_references += sum(raw.count(token) for token in legacy_tokens)

    report: dict[str, object] = {
        "manifest_entries": len(entries),
        "permanent_aliases": len(entries),
        "direct_renames": 0,
        "quatrefoil_aliases": sum("quatrefoil" in entry.old_path for entry in entries),
        "aged_aliases": sum("Aged order" in entry.reason for entry in entries),
        "slice_aliases": sum("Slice order" in entry.reason for entry in entries),
        "cylinder_aged_aliases": sum("cylinder_" in entry.old_path for entry in entries),
        "held_column_round_renames": sum("column_round" in entry.canonical_path for entry in entries),
        "canonical_registry_resources": canonical_resources,
        "stale_registry_resources": stale_resources,
        "canonical_language_keys": language_keys,
        "language_coverage": language_coverage,
        "stale_language_keys": stale_language_keys,
        "stale_registry_references": stale_registry_references,
        "ctm_property_files": len(ctm_roots),
        "ctm_legacy_references_preserved": ctm_legacy_references,
        "ctm_property_boms": ctm_property_boms,
        "errors": errors,
    }
    return report


def main() -> int:
    arguments = parser().parse_args()
    try:
        if arguments.command == "export":
            entries = workbook_entries(arguments.workbook.resolve(), arguments.sheet)
            write_manifest(entries, arguments.output.resolve())
            print(json.dumps({
                "output": str(arguments.output.resolve()),
                "entries": len(entries),
                "aliases": len(entries),
                "direct_renames": 0,
            }, indent=2, ensure_ascii=False))
            return 0

        entries = read_manifest(arguments.manifest.resolve())
        if arguments.command == "apply":
            print(json.dumps(apply_source(arguments.repo, entries), indent=2))
            return 0

        report = audit_source(arguments.repo, entries)
        if arguments.report:
            arguments.report.parent.mkdir(parents=True, exist_ok=True)
            arguments.report.write_text(
                json.dumps(report, indent=2, ensure_ascii=False) + "\n",
                encoding="utf-8")
        print(json.dumps(report, indent=2, ensure_ascii=False))
        return 1 if report["errors"] else 0
    except (OSError, ValueError, RuntimeError) as exception:
        print(f"id migration error: {exception}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
